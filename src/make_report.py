import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
R = "reports/"
d = Document(); d.styles["Normal"].font.name = "Calibri"; d.styles["Normal"].font.size = Pt(11)


def table(df, cols, names):
    t = d.add_table(rows=1, cols=len(cols)); t.style = "Light Grid Accent 1"
    for i, n in enumerate(names): t.rows[0].cells[i].text = n
    for _, r in df.iterrows():
        cells = t.add_row().cells
        for i, c in enumerate(cols): cells[i].text = f"{r[c]:.3f}" if isinstance(r[c], float) else str(r[c])


def H(s, l=1): d.add_heading(s, l)
def P(s): d.add_paragraph(s)
def B(s): d.add_paragraph(s, style="List Bullet")
def F(img, cap): d.add_picture(R + img, width=Inches(6.2)); d.add_paragraph(cap).italic = True


d.add_heading("Makine Öğrenmesi ile Kalite Tahmini: SECOM Yarı İletken Üretim Verisi Vaka Çalışması", 0)
P("[Ad Soyad – Öğrenci No] · [Arkadaşın Adı – Öğrenci No] · [Ders Adı] · [Tarih]")

H("1. Giriş")
P("Kestirimci kalite (predictive quality), üretim sırasında toplanan sensör ve süreç verilerinden ürünün kusurlu olup olmayacağını nihai testten önce tahmin etmeyi amaçlar. Kusurlu ürünün erken tespiti hurda, yeniden işleme ve test maliyetlerini azaltabilir. Bu çalışmada açık kaynaklı SECOM veri seti üzerinde (i) farklı makine öğrenmesi modellerini karşılaştırmak, (ii) sonuçların zaman içinde genelleşip genelleşmediğini sınamak ve (iii) tahmin hatalarının maliyetini dikkate alan bir karar eşiği analizi yapmak hedeflenmiştir.")

H("2. Veri")
P("SECOM (UCI Machine Learning Repository), bir yarı iletken üretim hattından 19 Temmuz – 17 Ekim 2008 arasında toplanan 1567 gözlem ve 590 anonim sensör ölçümünden oluşur. Gözlemlerin %6,6'sı (104 adet) kusurlu olarak etiketlenmiştir; yani sınıflar belirgin biçimde dengesizdir. Verinin yaklaşık %4,5'i eksiktir; 28 sensörde eksik oranı %50'nin üzerindedir ve 116 sensör sabit değerlidir.")
F("eda.png", "Şekil 1. Sınıf dağılımı, sensör başına eksik oranı ve haftalık kusur oranı.")

H("3. Yöntem")
B("Ön işleme: %50'den fazla eksik sensörler çıkarıldı; kalanlarda medyan ile doldurma, sabit sütun eleme ve standardizasyon uygulandı.")
B("Özellik seçimi: karşılıklı bilgi (mutual information) ile en bilgilendirici 40 sensör.")
B("Dengesizlik: sınıf ağırlıklandırma ve SMOTE.")
B("Modeller: Lojistik Regresyon (baz model), Random Forest, XGBoost.")
B("Sızıntı önleme: doldurma, seçim, ölçekleme ve SMOTE adımlarının tamamı bir boru hattı (pipeline) içinde, yalnızca eğitim katmanında öğrenilmiştir.")
B("Doğrulama: (a) 5 katlı tabakalı çapraz doğrulama × 3 tekrar, (b) zaman sıralı %70/%30 ayrım, (c) walk-forward: 500 gözlemle başlayıp 150'lik bloklarla ileriye doğru tahmin.")
B("Metrikler: dengesiz sınıflar nedeniyle doğruluk (accuracy) yerine PR-AUC, ROC-AUC ve en iyi eşikte F1. Rastgele bir sınıflayıcının PR-AUC değeri kusur oranına eşittir (≈0,066).")

H("4. Bulgular")
H("4.1 Model karşılaştırması", 2)
m = pd.read_csv(R + "model_results.csv")
table(m, ["model", "cv_prauc", "cv_roc", "cv_bestF1", "time_prauc", "time_roc"], ["Model", "CV PR-AUC", "CV ROC-AUC", "CV en iyi F1", "Zaman PR-AUC", "Zaman ROC-AUC"])
P("")
P("Rastgele çapraz doğrulamada en iyi model Random Forest'tır (PR-AUC 0,187; rastgele tahminin yaklaşık 2,8 katı). Ancak zaman sıralı ayrımda (son %30 test; baz oran 0,055) tüm modellerin performansı rastgele düzeye yaklaşmaktadır.")
H("4.2 Zamansal genelleme (walk-forward)", 2)
t = pd.read_csv(R + "temporal_results.csv")
table(t, ["config", "prauc", "roc", "base"], ["Konfigürasyon", "PR-AUC", "ROC-AUC", "Baz oran"])
P("")
P("Walk-forward değerlendirmede en iyi sonuç tüm sensörleri kullanan Random Forest'ta (PR-AUC 0,067) elde edilmiş, bu değer baz oranın (0,042) yaklaşık 1,6 katıdır. Özellik sayısını azaltmak veya kayan pencere kullanmak sonucu iyileştirmemiştir. Bu durum, mutual information ile seçilen sensörlerin eğitim dönemine özgü örüntüler taşıdığına ve sürecin zamanla değiştiğine (concept drift) işaret etmektedir. Rastgele çapraz doğrulamadaki iyimser sonuç, zaman sıralı doğrulama olmadan bu modelin değerinin olduğundan yüksek algılanacağını göstermektedir.")

H("4.3 Maliyet duyarlı karar eşiği", 2)
P("Kaçırılan kusurlu ürün maliyeti (C_FN) ile gereksiz yere ayıklanan sağlam ürün maliyeti (C_FP = 1 birim) için gerçek değerler veri setinde bulunmadığından, farklı oranlar senaryo olarak değerlendirilmiştir. Random Forest'ın çapraz doğrulama tahminleri üzerinde toplam maliyeti en aza indiren eşik hesaplanmıştır.")
c = pd.DataFrame({"Oran C_FN/C_FP": [5, 10, 20, 50], "En iyi eşik": [0.25, 0.11, 0.09, 0.09], "Model maliyeti": [466, 672, 852, 1212], "Hiç kontrol yok": [520, 1040, 2080, 5200], "Hepsini ayıkla": [1463] * 4})
table(c, list(c.columns), list(c.columns))
P("")
P("Kaçırma maliyeti yükseldikçe model tabanlı eşik, her iki uç stratejiye göre daha düşük maliyet sağlamaktadır; oran 10'da kazanç kontrolsüz stratejiye göre yaklaşık %35'tir. Bu sayılar çapraz doğrulama tahminlerine dayanır ve yukarıdaki zamansal sınırlılık nedeniyle iyimser olabilir.")
F("cost_threshold.png", "Şekil 2. Eşiğe göre toplam maliyet (Random Forest, çapraz doğrulama tahminleri).")

H("4.4 Yorumlama", 2)
P("SHAP analizine göre kusur tahminine en çok katkı veren sensörler s477, s33, s510, s65 ve s130'dur. Sensörler anonim olduğundan fiziksel yorum yapılamamış, ancak gerçek bir uygulamada bu sensörler süreç mühendisleriyle birlikte kök neden analizine başlangıç noktası olabilir.")
F("shap.png", "Şekil 3. SHAP özet grafiği (Random Forest).")

H("5. Sınırlılıklar")
B("Yalnızca 104 kusurlu gözlem vardır; tahminlerin belirsizliği yüksektir.")
B("Sensörler anonimdir; alan bilgisiyle özellik üretimi yapılamamıştır.")
B("Zamansal kayma nedeniyle model ileriye genelleşmemektedir.")
B("Maliyet oranları varsayımdır.")
B("Hiperparametre ayarı bilinçli olarak sınırlı tutulmuştur; az sayıda pozitif örnekle ayar, gürültüye aşırı uyum riski taşır.")

H("6. Sonuç ve Öneriler")
P("Çalışma, SECOM verisinde rastgele çapraz doğrulamayla umut verici görünen sonuçların zaman sıralı doğrulamada büyük ölçüde kaybolduğunu göstermiştir. Kestirimci kalite uygulamalarında doğrulama tasarımının zamanı dikkate alması gerektiği, bu çalışmanın ana çıkarımıdır. Maliyet duyarlı eşik analizi ise zayıf bir modelin bile doğru karar kuralıyla ekonomik değer üretebileceğini göstermektedir. Gelecek çalışmalarda drift tespiti, süreç kontrol kartlarıyla (SPC) hibrit kullanım ve anlamları bilinen sensörlerle özellik mühendisliği önerilir.")

H("Kaynaklar")
P("McCann, M. & Johnston, A. (2008). SECOM Data Set. UCI Machine Learning Repository.")
d.save(R + "Rapor_SECOM.docx")
