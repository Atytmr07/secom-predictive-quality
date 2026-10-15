from docx import Document
from docx.shared import Pt, Inches

d = Document()
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
H = lambda s, l=1: d.add_heading(s, l)
P = lambda s: d.add_paragraph(s)
B = lambda s: d.add_paragraph(s, style="List Bullet")


def table(header, rows, size=9.5):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
    for r in rows:
        c = t.add_row().cells
        for i, v in enumerate(r):
            c[i].text = v
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(size)
    d.add_paragraph("")


d.add_heading("Report 2: Methodology and Research Design", 0)
d.add_heading("Predictive Quality in Semiconductor Manufacturing: Can Defect-Prediction Models Survive Process Drift?", 2)
P("[Student Name 1 – ID] · [Student Name 2 – ID]\nIndependent Study and Industrial Applications · Due: 13 November 2026\n"
  "DRAFT – to be revised after Report 1 feedback. Items in [brackets] must be completed by the authors.")

# ------------------------------------------------------------------ 1
H("1. Refined Research Problem")
P("Report 1 identified that studies on the SECOM semiconductor dataset concentrate on data preparation and on improving prediction scores, while the concept-drift literature suggests that models validated on randomly selected records may be over-optimistic about future production. [Revise this paragraph after Report 1 feedback.]")
P("To keep the study achievable within the semester, the scope is narrowed to one public dataset (SECOM), three standard classifiers, one classical multivariate SPC baseline, and four validation protocols. Deep learning, online/adaptive learning and additional datasets are explicitly left to future work.")
P("Transparency note. A pilot version of the analysis was run while preparing Report 1 to check feasibility (code: https://github.com/Atytmr07/secom-predictive-quality). The hypotheses below were therefore formulated with some knowledge of the pilot. In Report 3, results will be reported as exploratory where this matters, and the full analysis will be re-run with fixed random seeds before any value is reported.")

# ------------------------------------------------------------------ 2
H("2. Final Research Objective, Research Questions and Hypotheses")
H("Research objective", 2)
P("To measure how much the apparent performance of defect-prediction models on the SECOM dataset depends on the validation protocol, to quantify sensor drift and test whether it explains the loss of performance on later production runs, and to compare machine-learning models with a PCA-based Hotelling T² control chart using statistical and cost-based criteria.")
H("Research questions and hypotheses", 2)
table(["RQ", "Question", "Hypothesis", "Analysis that answers it"], [
    ("RQ1", "How does the performance of Logistic Regression, Random Forest and XGBoost differ between random cross-validation and time-aware validation?",
     "H1: For each model, lift over random guessing is lower under chronological and walk-forward validation than under random stratified cross-validation.",
     "Protocol comparison (Section 7.1) with bootstrap confidence intervals of lift"),
    ("RQ2", "How much do sensor distributions change over time, and is the loss of performance related to drift?",
     "H2a: Sensors flagged as drifting within the training period shift more between training and test periods than stable sensors. H2b: Models trained only on stable sensors perform better on later runs than models trained only on drifting sensors of the same number.",
     "KS-test drift analysis and stable/drifting/random sensor-subset experiment (Section 7.2)"),
    ("RQ3", "How does a PCA-based Hotelling T² control chart compare with machine-learning models on later production runs?",
     "H3: On chronologically held-out runs, the PR-AUC of the best ML model is not significantly higher than that of the T² chart.",
     "Chronological comparison with bootstrap confidence intervals (Section 7.3)"),
    ("RQ4", "Under assumed cost ratios between missed defects and false alarms, which threshold minimizes total cost, and does a model beat no inspection and full inspection?",
     "H4: For sufficiently high missed-defect/false-alarm cost ratios, a model-based threshold has lower total cost than both trivial strategies in cross-validation; this advantage is expected to be smaller or absent under chronological validation (threshold chosen on the training period only).",
     "Cost-threshold analysis for several cost ratios under both protocols (Section 7.4)"),
])
P("Hypotheses are tested descriptively with confidence intervals rather than with a single p-value, because the number of defective runs (104) is small and the evaluation metrics are not normally distributed.")

# ------------------------------------------------------------------ 3
H("3. Research Methodology")
P("The study is a quantitative, data-driven experimental comparison built on supervised classification, multivariate statistical process control and distribution-shift testing. Each choice is justified below.")
table(["Component", "Choice", "Justification"], [
    ("Task", "Binary classification (fail = 1, pass = 0)", "The final test gives a pass/fail label; predictive quality in this setting is a classification problem (Tercan & Meisen, 2022)."),
    ("Models", "Logistic Regression (class-weighted), Random Forest, XGBoost", "A transparent linear baseline and two tree ensembles that are widely used for tabular manufacturing data and handle many correlated, partly missing features. Using standard models keeps the focus on validation, not on model novelty."),
    ("Imbalance handling", "Class weights; SMOTE as a comparison for Logistic Regression", "About 6.6% of runs are defective. Class weighting does not change the data; SMOTE is included because most SECOM studies use resampling (e.g. Park et al., 2024; Nuhu et al., 2023)."),
    ("Preprocessing", "Remove sensors with >50% missing values and constant sensors; median imputation; scaling; mutual-information selection of 40 sensors", "Simple, transparent steps. All steps are fitted only on training data inside each fold (a pipeline) to avoid data leakage."),
    ("SPC baseline", "PCA (90% variance) on good runs of the training period; Hotelling T² statistic as anomaly score; upper control limit = 99th percentile of training T²", "T² with projection methods is the standard multivariate SPC tool for many correlated variables (Kourti & MacGregor, 1995) and needs only normal-operation data."),
    ("Drift test", "Two-sample Kolmogorov–Smirnov test per sensor, Bonferroni-corrected", "Non-parametric, no distributional assumption; applicable to every sensor. Drifting sensors are defined using the training period only (first half vs second half) to prevent information from the test period leaking into the selection."),
    ("Validation protocols", "Random stratified 5-fold; blocked (contiguous) 5-fold; chronological 70/30 split; walk-forward with expanding window", "These form a gradient from ignoring time to respecting it fully. Comparing them directly answers RQ1 (Gama et al., 2014; Jourdan et al., 2021)."),
    ("Uncertainty", "Bootstrap 95% confidence intervals", "With only 104 positives, point estimates are unstable; intervals show whether a result differs from random guessing."),
])

# ------------------------------------------------------------------ 4
H("4. Research Framework / Methodological Flowchart")
d.add_picture("reports/flowchart.png", width=Inches(4.6))
P("Figure 1. Methodological flowchart of the study.")

# ------------------------------------------------------------------ 5
H("5. Dataset and Data Collection")
table(["Item", "Description"], [
    ("Source", "UCI Machine Learning Repository, SECOM dataset (McCann & Johnston, 2008), https://archive.ics.uci.edu/dataset/179/secom. Data were downloaded on 5 October 2026; no new data are collected."),
    ("Observations", "1,567 production runs from a semiconductor manufacturing process."),
    ("Time period", "19 July 2008 to 17 October 2008 (timestamps provided with the labels)."),
    ("Target", "Pass/fail label from the in-house test. 1,463 pass and 104 fail (6.6%). The original coding (-1 = pass, 1 = fail) is recoded to 0/1."),
    ("Features", "590 anonymised numeric sensor measurements. Physical meaning, units and sensor type are not disclosed."),
    ("Missing values", "About 4.5% of all cells are missing. 28 sensors have more than 50% missing values and are removed; among the remaining sensors the median missing rate is about 0.2% and the maximum about 46%."),
    ("Constant sensors", "116 sensors have a single value and are removed. After removing constant and sparse sensors, 446 sensors remain."),
    ("Outliers", "Using a robust z-score (median absolute deviation) threshold of 3.5, about 3% of all remaining cells are flagged; 74 sensors have more than 5% flagged values. Outliers are kept, because in a defect-detection setting extreme values may carry the signal. Defective runs have on average more flagged sensors (16.2) than good runs (13.2); this is descriptive."),
    ("Correlation structure", "347 sensor pairs have an absolute correlation above 0.9, i.e. many sensors are redundant. This motivates feature selection and PCA."),
    ("Data quality issues", "Anonymous sensors (no domain feature engineering), few positives, strong imbalance, possible process changes during the three months."),
])
P("The figures in this section were computed from the data with src/eda_extra.py and src/eda.py in the project repository.")

# ------------------------------------------------------------------ 6
H("6. Variables")
table(["Variable", "Description", "Type", "Unit", "Role in analysis"], [
    ("S1 … S590 (s0 … s589 in code)", "Anonymised process/sensor measurements", "Numerical (continuous)", "Not disclosed", "Independent (features); 446 retained after cleaning"),
    ("Time stamp", "Date and time of the production run", "Date-time", "Date-time", "Used only to order runs and to define chronological splits and drift periods; not used as a feature"),
    ("Y (label)", "Final in-house test result: fail = 1, pass = 0", "Binary categorical", "–", "Dependent (target)"),
    ("T² score", "Hotelling T² statistic of a run in the PCA space of good runs", "Numerical (derived)", "Dimensionless", "Anomaly score of the SPC baseline"),
    ("Model score", "Predicted probability of failure from a classifier", "Numerical (derived)", "0–1", "Used for threshold selection and evaluation"),
    ("Drift status", "Whether a sensor differs significantly between the first and second half of the training period", "Binary (derived)", "–", "Grouping variable for the stable/drifting sensor experiment"),
])

# ------------------------------------------------------------------ 7
H("7. Proposed Analysis")
H("7.1 Validation-protocol comparison (RQ1)", 2)
P("The same three models and the same preprocessing pipeline are evaluated under four protocols: random stratified 5-fold CV, blocked 5-fold CV without shuffling, a single chronological 70/30 split, and walk-forward evaluation (initial 500 runs, 150-run blocks, expanding window). Test predictions are pooled and PR-AUC, lift (PR-AUC divided by the fraction of defects in the evaluated runs) and ROC-AUC are reported with bootstrap confidence intervals. Lift is used because the share of defects differs between protocols, so raw PR-AUC values are not comparable across them.")
H("7.2 Drift analysis and its link to performance (RQ2)", 2)
P("For every retained sensor a two-sample KS test compares the distribution in the first and the last part of the data. To test whether drift explains the performance loss without leaking information, sensors are labelled stable or drifting using only the first half versus the second half of the training period (first 70%). Models are then trained on equal-sized random subsets of stable sensors, of drifting sensors, and of all sensors (20 repetitions each) and evaluated on the last 30%. The experiment also checks whether sensors flagged within the training period do shift more between training and test periods.")
H("7.3 SPC versus machine learning (RQ3)", 2)
P("A PCA model is fitted on good runs of the training period only, T² is computed for later runs, and its PR-AUC and ROC-AUC on the chronologically held-out runs are compared with the machine-learning models, using bootstrap confidence intervals.")
H("7.4 Cost-based decision analysis (RQ4)", 2)
P("Total cost = C_FN × (missed defects) + C_FP × (good runs flagged), with C_FP = 1 and C_FN/C_FP in {5, 10, 20, 50}. No real cost data exist for this dataset, so the ratios are assumptions and the analysis is a sensitivity analysis, not an estimate of real savings. For each ratio, the threshold that minimizes total cost is found and compared with 'no inspection' (all runs accepted) and 'full inspection' (all runs flagged). The analysis is performed on cross-validated predictions and, as a stricter check, on chronologically held-out predictions.")
H("7.5 Interpretation", 2)
P("SHAP values of the Random Forest are used to see which sensors contribute most. Because sensors are anonymous, this analysis is descriptive and no physical interpretation is claimed.")

# ------------------------------------------------------------------ 8
H("8. Evaluation Criteria")
table(["Criterion", "Why it is appropriate here"], [
    ("PR-AUC (average precision)", "Summarizes the precision–recall trade-off over all thresholds and focuses on the rare defective class. Its random-guess value equals the defect rate, which makes interpretation easy."),
    ("Lift = PR-AUC / defect rate", "Makes results comparable across protocols whose test sets have different defect rates."),
    ("ROC-AUC", "Threshold-independent ranking quality; reported as a secondary measure because it can look good under strong imbalance."),
    ("Best-threshold F1 and MCC", "Single-number summaries of a chosen decision rule; MCC uses all four cells of the confusion matrix and is robust to imbalance. Used descriptively only."),
    ("Total cost under assumed cost ratios", "Industrial Engineering decisions depend on the cost of wrong decisions; a missed defect and a false alarm are not equally costly."),
    ("Bootstrap 95% confidence interval", "With 104 positives, intervals show whether the difference from random guessing or between methods is meaningful."),
])
P("Accuracy is not used as a performance measure: predicting 'pass' for every run would already give about 93% accuracy while detecting no defects.")

# ------------------------------------------------------------------ 9
H("9. Expected Results")
P("Expected results are expectations, not findings. Based on the literature in Report 1 we expect: (i) performance to be highest under random cross-validation and lower under time-aware protocols; (ii) a large share of sensors to show statistically significant shifts over three months; (iii) the classical T² chart to be a reasonable but not clearly better or worse alternative to machine learning on later runs; and (iv) a model-based threshold to be more useful when missed defects are much more costly than false alarms. We do not assume any particular numerical value. If the results contradict these expectations, they will be reported as they are.")

H("Limitations of the Design")
B("Single dataset from one process and period; the findings may not generalize to other factories.")
B("Only 104 defective runs, so confidence intervals will be wide.")
B("Anonymous sensors prevent domain-based feature engineering and physical interpretation.")
B("Cost ratios are assumed; real plant costs are not available.")
B("Hyperparameters are mostly kept at reasonable defaults; extensive tuning with so few positives risks overfitting to noise.")

H("References")
refs = [
    "Gama, J., Žliobaitė, I., Bifet, A., Pechenizkiy, M., & Bouchachia, A. (2014). A survey on concept drift adaptation. ACM Computing Surveys, 46(4), Article 44. https://doi.org/10.1145/2523813",
    "Jourdan, N., Sen, S., Husom, E. J., Garcia-Ceja, E., Biegel, T., & Metternich, J. (2021). On the reliability of machine learning applications in manufacturing environments. arXiv:2112.06986 (preprint).",
    "Kourti, T., & MacGregor, J. F. (1995). Process analysis, monitoring and diagnosis, using multivariate projection methods. Chemometrics and Intelligent Laboratory Systems, 28, 3–21. https://doi.org/10.1016/0169-7439(95)80036-9",
    "McCann, M., & Johnston, A. (2008). SECOM [Dataset]. UCI Machine Learning Repository. https://archive.ics.uci.edu/dataset/179/secom",
    "Nuhu, A. A., Zeeshan, Q., Safaei, B., & Shahzad, M. A. (2023). Machine learning-based techniques for fault diagnosis in the semiconductor manufacturing process: a comparative study. The Journal of Supercomputing.",
    "Park, H.-J., Koo, Y.-S., Yang, H.-Y., Han, Y.-S., & Nam, C.-S. (2024). Study on data preprocessing for machine learning based on semiconductor manufacturing processes. Sensors, 24(17), 5461. https://doi.org/10.3390/s24175461",
    "Tercan, H., & Meisen, T. (2022). Machine learning and deep learning based predictive quality in manufacturing: a systematic review. Journal of Intelligent Manufacturing, 33(7), 1879–1905.",
]
for r in refs:
    P(r)
d.save("reports/Report2_Methodology.docx")
